import os, re, sys
import pandas as pd
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

M = '******'
# Payer/employer names to mask on credit SMS. The real pattern is private: set LEDGER_EMPLOYER_RE.
EMPLOYER_RE = os.environ.get('LEDGER_EMPLOYER_RE', r'EMPLOYER NAME')
CUT = pd.Timestamp('2025-12-31 23:59:59.999')
NOT_NAMES = r'(?:Customer|Customers|Valued|Cardholder|FAB Customer|there|There|Team)\b'
UTILITY_NICKNAMES = re.compile(r'^(?:HomeDewa|HomeELife|Du|Etisalat|DEWA|SEWA)\b', re.I)

def mask_beneficiary(m):
    name = m.group(2)
    return m.group(0) if UTILITY_NICKNAMES.match(name.strip()) else m.group(1) + M + m.group(3)

RULES = [
    # People's names
    ('greeting name', r'\b(Dear|Hi|Hello)\s+(?!' + NOT_NAMES + r')[A-Za-z][A-Za-z .]*?(?=\s*[,!])', r'\1 ' + M),
    ('beneficiary name', r'(added )(.+?)( as a new beneficiary)', mask_beneficiary),
    ('beneficiary name', r'(newly added beneficiary: )(.+?)( has been activated)', mask_beneficiary),
    ('beneficiary name', r'(beneficiary creation/modification request for )(.+?)( XXXX\d+| -)', mask_beneficiary),
    ('transfer recipient', r'(Your transfer to )(.+?)( for AED)', r'\1' + M + r'\3'),
    ('transfer recipient', r'(transfer to )(?!IBAN|your)([A-Z][^.]*?)(\. If this is not you)', r'\1' + M + r'\3'),
    ('transfer recipient', r'(transfer request \(Ref No \S+\) for [A-Z]{3} [\d,.]+ to )([^.]+)', r'\1' + M),
    # Payers, employers and transfer notes on credits
    ('payer / employer', r'(\bB(?:Y|/O) )(?:' + EMPLOYER_RE + r').*?(?=\. The available|$)', r'\1' + M),
    ('payer / employer', r'(B/O ?)(?:' + EMPLOYER_RE + r').*?(?=\. The available|$)', r'\1' + M),
    ('payer / employer', r'(IPI TT REF: \S+ ).*?(?=\. The available|$)', r'\1' + M),
    ('transfer memo', r'(;MSG-)(?!\*+(?:\.|$))[^.\n]*', r'\1' + M),
    # Account, card, IBAN and customer numbers
    ('IBAN', r'\bAE\d{2}[\d ]{10,}\d|\bAE\d+XXX[\dX]+', 'AE' + M),
    ('full card number', r'\b\d{4}[x*]{6,}(\d{4})\b', r'XXXX\1'),
    ('utility account no', r'\b((?:DU|DEWA|SEWA|ETISALAT|FEWA|ADDC) NO\.-)\d+', r'\1' + M),
    ('utility account no', r'(consumer number )\d+', r'\1' + M),
    ('reference / complaint no', r'((?:reference (?:number|no\.?)|request|complaint|Request ID:|application reference is|customer reference number is)\s*#?\s*)(L-)?[A-Z]*\d{5,}', r'\1' + M),
    ('cheque number', r'(cheque (?:number|no\.?:?)|Cheque No\.?)\s*[X\d]+', r'\1 ' + M),
    ('OTP / code', r'((?:OTP|code|password)(?: is|:)?\s*(?:CODE TO ANYONE\.\s*)?)\d{4,8}', r'\1' + M),
    ('OTP / code', r'\b\d{4,8}(?= is the OTP| is your OTP| is OTP)', M),
    ('OTP / code', r'(\b(?:OTP|One Time Password|code)\b[^.\d]{0,80}?\bis:?\s*)\d{4,8}', r'\1' + M),
    ('OTP / code', r'(enter this OTP )\d+', r'\1' + M),
    ('employer as merchant', r'(at )(?:' + EMPLOYER_RE + r')[^,.]*', r'\1' + M),
    ('email', r'[\w.+-]+@(?!google\.com|liv\.me)[\w-]+\.[\w.]+', M),
]

def censor(text):
    hits = []
    for name, pat, rep in RULES:
        new = re.sub(pat, rep, text, flags=re.IGNORECASE if name in ('OTP / code', 'cheque number', 'reference / complaint no') else 0)
        if new != text:
            hits.append(name); text = new
    return text, hits

def write_xlsx(df, path, sheet='SMS'):
    wb = Workbook(); ws = wb.active; ws.title = sheet
    ws.append(['Date & Time (GST)', 'Sender', 'Message', 'Read'])
    for r in df.itertuples(index=False):
        ws.append([r.date.to_pydatetime(), r.sender, r.message, 'Yes' if r.read == '1' else 'No'])
    f = Font(name='Arial', size=10)
    for row in ws.iter_rows():
        for c in row: c.font = f; c.alignment = Alignment(vertical='top', wrap_text=(c.column == 3))
    for c in ws[1]: c.font = Font(name='Arial', size=10, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78')
    for c in ws['A'][1:]: c.number_format = 'dd mmm yyyy hh:mm:ss'
    for col, w in zip('ABCD', [20, 10, 90, 7]): ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    wb.save(path)

if __name__ == '__main__':
    for name, out in [('liv', 'contributor-c-liv-redacted-2025.xlsx'), ('fab', 'contributor-d-fab-redacted-2025.xlsx')]:
        d = pd.read_pickle(f'{name}_raw.pkl')
        n0 = len(d); d = d[d.date <= CUT].copy()
        counts = Counter(); msgs = []
        for t in d.message:
            t2, h = censor(t); counts.update(set(h)); msgs.append(t2)
        d['message'] = msgs
        d.to_pickle(f'{name}_censored.pkl')
        write_xlsx(d, out)
        print(out, n0, '->', len(d), d.date.min(), d.date.max(), dict(counts))
